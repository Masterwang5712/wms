"""Excel 导出（openpyxl），统一表头样式、冻结首行、自动列宽、金额格式化。"""
from datetime import date
from io import BytesIO
from urllib.parse import quote

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..config import CALIBRATION_WARN_DAYS, DEVICE_STATUS
from ..db import get_conn, query_all
from ..security import current_user
from ..utils import month_range, days_between, round2, today_str

router = APIRouter(prefix="/export", tags=["导出"])

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

HEADER_FILL = PatternFill("solid", fgColor="2F6FED")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
THIN = Side(style="thin", color="D9DDE6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")
LOW_FILL = PatternFill("solid", fgColor="FFF1E0")
OUT_FILL = PatternFill("solid", fgColor="FDE2E2")
EXPIRED_FILL = PatternFill("solid", fgColor="FDE2E2")
EXPIRING_FILL = PatternFill("solid", fgColor="FFF6E0")
MONEY_FMT = '¥#,##0.00'


def _sheet(wb, title: str, headers: list, rows: list, *, money_cols=(), int_cols=(),
           widths=None, row_fill=None, first=False):
    """创建并填充工作表。first=True 时复用 Workbook 自带的默认表，避免残留空白 Sheet。"""
    safe_title = title[:31]
    if first:
        ws = wb.active
        ws.title = safe_title
    else:
        ws = wb.create_sheet(title=safe_title)

    ws.append(headers)
    for ci, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=ci)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[1].height = 22

    for ri, r in enumerate(rows, start=2):
        ws.append(list(r))
        fill = row_fill(r) if row_fill else None
        for ci in range(1, len(headers) + 1):
            cell = ws.cell(row=ri, column=ci)
            cell.border = BORDER
            if ci in money_cols:
                cell.number_format = MONEY_FMT
                cell.alignment = RIGHT
            elif ci in int_cols:
                cell.alignment = RIGHT
            else:
                cell.alignment = LEFT
            if fill:
                cell.fill = fill

    # 列宽
    for ci in range(1, len(headers) + 1):
        if widths and ci - 1 < len(widths):
            w = widths[ci - 1]
        else:
            longest = len(str(headers[ci - 1]))
            for r in rows[:200]:
                v = r[ci - 1] if ci - 1 < len(r) else ""
                longest = max(longest, len(str(v if v is not None else "")))
            w = min(max(longest + 4, 10), 42)
        ws.column_dimensions[get_column_letter(ci)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    return ws


def _stream(wb: Workbook, filename: str) -> StreamingResponse:
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    # RFC5987 编码中文文件名
    disposition = f"attachment; filename=\"export.xlsx\"; filename*=UTF-8''{quote(filename)}"
    return StreamingResponse(bio, media_type=XLSX_MIME, headers={"Content-Disposition": disposition})


def _file_name(base: str) -> str:
    return f"{base}_{date.today().strftime('%Y%m%d')}.xlsx"


@router.get("/inventory.xlsx")
def export_inventory(conn=Depends(get_conn), user: dict = Depends(current_user)):
    rows = query_all(conn, """
        SELECT i.code, i.name, c.name AS category_name, i.spec, i.unit, i.brand, i.location,
               i.current_stock, i.safety_stock, i.unit_price,
               ROUND(i.current_stock * i.unit_price, 2) AS stock_value,
               CASE WHEN i.current_stock<=0 THEN '缺货'
                    WHEN i.safety_stock>0 AND i.current_stock<=i.safety_stock THEN '低库存'
                    ELSE '正常' END AS status
        FROM item i LEFT JOIN item_category c ON c.id=i.category_id
        WHERE i.is_deleted=0 AND i.item_type='consumable' ORDER BY i.code
    """)

    def fill(r):
        if r[7] <= 0:
            return OUT_FILL
        if r[8] > 0 and r[7] <= r[8]:
            return LOW_FILL
        return None

    wb = Workbook()
    _sheet(wb, "实时库存",
           ["物品编码", "物品名称", "分类", "规格", "单位", "品牌", "存放位置",
            "当前库存", "安全库存", "参考单价", "库存货值", "状态"],
           [list(r.values()) for r in rows],
           money_cols=(10, 11), int_cols=(8, 9), row_fill=fill, first=True,
           widths=[14, 26, 12, 24, 8, 12, 14, 10, 10, 12, 14, 10])
    total = query_all(conn, """
        SELECT COUNT(*), COALESCE(SUM(current_stock),0), ROUND(COALESCE(SUM(current_stock*unit_price),0),2)
        FROM item WHERE is_deleted=0 AND item_type='consumable'
    """)[0]
    tv = list(total.values())
    ws = wb["实时库存"]
    r = ws.max_row + 2
    ws.cell(row=r, column=1, value="合计").font = Font(bold=True)
    ws.cell(row=r, column=8, value=f"品类 {tv[0]} 种 · 库存合计 {tv[1]} 件").font = Font(bold=True)
    cell = ws.cell(row=r, column=11, value=round2(tv[2]))
    cell.font = Font(bold=True)
    cell.number_format = MONEY_FMT
    return _stream(wb, _file_name("实时库存表"))


@router.get("/items.xlsx")
def export_items(conn=Depends(get_conn), user: dict = Depends(current_user)):
    rows = query_all(conn, """
        SELECT i.code, i.name, c.name AS category_name,
               CASE i.item_type WHEN 'consumable' THEN '耗材' ELSE '设备' END AS type_label,
               i.spec, i.unit, i.brand, i.safety_stock, i.current_stock, i.unit_price, i.location, i.remark
        FROM item i LEFT JOIN item_category c ON c.id=i.category_id
        WHERE i.is_deleted=0 ORDER BY i.code
    """)
    wb = Workbook()
    _sheet(wb, "物品档案",
           ["物品编码", "物品名称", "分类", "类型", "规格", "单位", "品牌",
            "安全库存", "当前库存", "参考单价", "存放位置", "备注"],
           [list(r.values()) for r in rows], money_cols=(10,), int_cols=(8, 9), first=True)
    return _stream(wb, _file_name("物品档案表"))


@router.get("/txns.xlsx")
def export_txns(date_from: str | None = None, date_to: str | None = None,
                conn=Depends(get_conn), user: dict = Depends(current_user)):
    sql = """SELECT t.created_at, t.txn_type, i.code, i.name AS item_name, i.spec, i.unit,
                    t.quantity, t.unit_price, t.amount, t.before_qty, t.after_qty,
                    COALESCE(d.name,'-') AS department_name,
                    COALESCE(u.display_name,'-') AS operator_name, t.remark
             FROM stock_txn t
             JOIN item i ON i.id=t.item_id
             LEFT JOIN org_department d ON d.id=t.department_id
             LEFT JOIN sys_user u ON u.id=t.operator_id
             WHERE t.txn_type IN ('in','out')"""
    params = []
    if date_from:
        sql += " AND date(t.created_at) >= date(?)"
        params.append(date_from)
    if date_to:
        sql += " AND date(t.created_at) <= date(?)"
        params.append(date_to)
    sql += " ORDER BY t.id DESC"
    rows = query_all(conn, sql, params)

    label = {"in": "入库", "out": "出库"}
    data = []
    for r in rows:
        v = list(r.values())
        v[1] = label.get(v[1], v[1])
        data.append(v)

    wb = Workbook()
    _sheet(wb, "出入库流水",
           ["时间", "类型", "物品编码", "物品名称", "规格", "单位", "数量", "单价", "金额",
            "变更前库存", "变更后库存", "领用部门", "操作人", "备注"],
           data, money_cols=(8, 9), int_cols=(7, 10, 11), first=True,
           widths=[20, 8, 14, 26, 22, 8, 10, 12, 14, 12, 12, 14, 14, 26])
    return _stream(wb, _file_name("出入库流水"))


@router.get("/consumption.xlsx")
def export_consumption(dim: str = "department", date_from: str | None = None,
                       date_to: str | None = None,
                       conn=Depends(get_conn), user: dict = Depends(current_user)):
    dim_map = {
        "department": ("COALESCE(d.name,'未指定部门')", "LEFT JOIN org_department d ON d.id=ci.department_id", "部门"),
        "staff": ("COALESCE(s.name,'未指定人员')", "LEFT JOIN org_staff s ON s.id=ci.staff_id", "人员"),
        "item": ("i.name || CASE WHEN i.spec IS NULL OR i.spec='' THEN '' ELSE ' / ' || i.spec END",
                 "JOIN item i ON i.id=ci.item_id", "物品"),
        "month": ("substr(ci.issued_at,1,7)", "", "月份"),
    }
    expr, join, label = dim_map.get(dim, dim_map["department"])
    sql = f"""SELECT {expr} AS name, SUM(ci.quantity) AS qty, ROUND(SUM(ci.amount),2) AS amount, COUNT(*) AS cnt
              FROM consumable_issue ci {join} WHERE 1=1"""
    params = []
    if date_from:
        sql += " AND date(ci.issued_at) >= date(?)"
        params.append(date_from)
    if date_to:
        sql += " AND date(ci.issued_at) <= date(?)"
        params.append(date_to)
    sql += " GROUP BY name ORDER BY amount DESC"
    rows = query_all(conn, sql, params)

    wb = Workbook()
    _sheet(wb, f"消耗统计-{label}", [label, "领用数量", "领用金额", "领用次数"],
           [list(r.values()) for r in rows], money_cols=(3,), int_cols=(2, 4), first=True,
           widths=[34, 14, 16, 12])
    return _stream(wb, _file_name(f"耗材消耗统计-{label}"))


@router.get("/devices.xlsx")
def export_devices(conn=Depends(get_conn), user: dict = Depends(current_user)):
    rows = query_all(conn, """
        SELECT dv.asset_no, dv.name, c.name AS category_name, dv.brand, dv.model, dv.serial_no,
               dv.purchase_date, dv.price, dv.location, s.name AS custodian_name,
               COALESCE(d.name,'-') AS department_name, dv.status,
               dv.calibration_cycle_days, dv.last_calibration_date, dv.next_calibration_date
        FROM device dv
        LEFT JOIN item_category c ON c.id=dv.category_id
        LEFT JOIN org_staff s ON s.id=dv.custodian_id
        LEFT JOIN org_department d ON d.id=dv.department_id
        WHERE dv.is_deleted=0 ORDER BY dv.asset_no
    """)
    data = []
    for r in rows:
        next_cal = r["next_calibration_date"]
        left = days_between(next_cal)
        if not next_cal:
            cal_status = "无需校准"
        elif left is None:
            cal_status = "-"
        elif left < 0:
            cal_status = f"已过期 {-left} 天"
        elif left <= CALIBRATION_WARN_DAYS:
            cal_status = f"临期（剩 {left} 天）"
        else:
            cal_status = f"正常（剩 {left} 天）"
        data.append([
            r["asset_no"], r["name"], r["category_name"], r["brand"], r["model"], r["serial_no"],
            r["purchase_date"], r["price"], r["location"], r["custodian_name"], r["department_name"],
            DEVICE_STATUS.get(r["status"], r["status"]),
            r["calibration_cycle_days"], r["last_calibration_date"], next_cal, cal_status,
        ])

    wb = Workbook()
    ws = _sheet(wb, "设备台账",
                ["资产编号", "设备名称", "分类", "品牌", "型号", "序列号", "购置日期", "购置金额",
                 "存放位置", "责任人", "使用部门", "状态", "校准周期(天)", "上次校准", "下次校准", "校准状态"],
                data, money_cols=(8,), int_cols=(13,), first=True,
                widths=[14, 26, 12, 12, 20, 20, 13, 14, 16, 12, 14, 10, 13, 13, 13, 16])
    # 按校准状态着色（第 16 列为校准状态）
    for ri, row in enumerate(data, start=2):
        st = str(row[15] or "")
        if st.startswith("已过期"):
            f = EXPIRED_FILL
        elif st.startswith("临期"):
            f = EXPIRING_FILL
        else:
            f = None
        if f:
            for ci in range(1, 17):
                ws.cell(row=ri, column=ci).fill = f
    return _stream(wb, _file_name("设备台账"))


@router.get("/monthly.xlsx")
def export_monthly(year: int | None = None, month: int | None = None,
                   conn=Depends(get_conn), user: dict = Depends(current_user)):
    t = date.today()
    year = year or t.year
    month = month or t.month
    start, end = month_range(year, month)

    wb = Workbook()

    # Sheet1 库存快照
    inv = query_all(conn, """
        SELECT i.code, i.name, c.name AS category_name, i.spec, i.unit,
               i.current_stock, i.safety_stock, i.unit_price,
               ROUND(i.current_stock*i.unit_price,2) AS value
        FROM item i LEFT JOIN item_category c ON c.id=i.category_id
        WHERE i.is_deleted=0 AND i.item_type='consumable' ORDER BY i.code
    """)
    _sheet(wb, "库存快照",
           ["物品编码", "物品名称", "分类", "规格", "单位", "当前库存", "安全库存", "单价", "库存货值"],
           [list(r.values()) for r in inv], money_cols=(8, 9), int_cols=(6, 7), first=True)

    # Sheet2 入库单
    inbound = query_all(conn, """
        SELECT o.order_no, o.order_date, o.supplier_name, o.handler, o.total_amount, o.status, o.remark
        FROM inbound_order o WHERE date(o.order_date) BETWEEN date(?) AND date(?) ORDER BY o.order_date
    """, (start, end))
    _sheet(wb, "入库单",
           ["入库单号", "入库日期", "供应商", "经办人", "金额", "状态", "备注"],
           [list(r.values()) for r in inbound], money_cols=(5,))

    # Sheet3 出入库流水
    txns = query_all(conn, """
        SELECT t.created_at, t.txn_type, i.code, i.name, i.unit, t.quantity, t.unit_price, t.amount,
               t.before_qty, t.after_qty, COALESCE(d.name,'-'), COALESCE(u.display_name,'-'), t.remark
        FROM stock_txn t JOIN item i ON i.id=t.item_id
        LEFT JOIN org_department d ON d.id=t.department_id
        LEFT JOIN sys_user u ON u.id=t.operator_id
        WHERE date(t.created_at) BETWEEN date(?) AND date(?) ORDER BY t.id
    """, (start, end))
    label = {"in": "入库", "out": "出库", "adjust": "盘点"}
    tx_data = []
    for r in txns:
        v = list(r.values())
        v[1] = label.get(v[1], v[1])
        tx_data.append(v)
    _sheet(wb, "出入库流水",
           ["时间", "类型", "物品编码", "物品名称", "单位", "数量", "单价", "金额", "变更前", "变更后", "部门", "操作人", "备注"],
           tx_data, money_cols=(7, 8), int_cols=(6, 9, 10))

    # Sheet4 部门消耗
    by_dept = query_all(conn, """
        SELECT COALESCE(d.name,'未指定') AS name, SUM(ci.quantity) AS qty,
               ROUND(SUM(ci.amount),2) AS amount, COUNT(*) AS cnt
        FROM consumable_issue ci LEFT JOIN org_department d ON d.id=ci.department_id
        WHERE date(ci.issued_at) BETWEEN date(?) AND date(?) GROUP BY d.name ORDER BY amount DESC
    """, (start, end))
    _sheet(wb, "部门消耗统计", ["部门", "领用数量", "领用金额", "领用次数"],
           [list(r.values()) for r in by_dept], money_cols=(3,), int_cols=(2, 4))

    # Sheet5 物品消耗排行
    by_item = query_all(conn, """
        SELECT i.code, i.name, i.spec, i.unit, SUM(ci.quantity) AS qty, ROUND(SUM(ci.amount),2) AS amount
        FROM consumable_issue ci JOIN item i ON i.id=ci.item_id
        WHERE date(ci.issued_at) BETWEEN date(?) AND date(?) GROUP BY ci.item_id ORDER BY amount DESC
    """, (start, end))
    _sheet(wb, "物品消耗排行", ["物品编码", "物品名称", "规格", "单位", "领用数量", "领用金额"],
           [list(r.values()) for r in by_item], money_cols=(6,))

    # Sheet6 维保与校准
    maint = query_all(conn, """
        SELECT m.maint_date, dv.asset_no, dv.name, m.maint_type, m.content, m.cost, m.vendor
        FROM device_maintenance m JOIN device dv ON dv.id=m.device_id
        WHERE date(m.maint_date) BETWEEN date(?) AND date(?) ORDER BY m.maint_date
    """, (start, end))
    calib = query_all(conn, """
        SELECT c.calibration_date, dv.asset_no, dv.name, c.result, c.agency,
               c.certificate_no, c.next_date, c.cost
        FROM device_calibration c JOIN device dv ON dv.id=c.device_id
        WHERE date(c.calibration_date) BETWEEN date(?) AND date(?) ORDER BY c.calibration_date
    """, (start, end))
    _sheet(wb, "维保记录", ["维保日期", "资产编号", "设备名称", "类型", "内容", "费用", "服务商"],
           [list(r.values()) for r in maint], money_cols=(6,))
    _sheet(wb, "校准记录", ["校准日期", "资产编号", "设备名称", "结论", "校准机构", "证书编号", "下次校准", "费用"],
           [list(r.values()) for r in calib], money_cols=(8,))

    return _stream(wb, f"仓库月度报表_{year}{month:02d}.xlsx")


@router.get("/calibration.xlsx")
def export_calibration(conn=Depends(get_conn), user: dict = Depends(current_user)):
    rows = query_all(conn, """
        SELECT dv.asset_no, dv.name, dv.model, COALESCE(d.name,'-') AS dept, dv.location,
               s.name AS custodian, dv.calibration_cycle_days, dv.last_calibration_date,
               dv.next_calibration_date
        FROM device dv
        LEFT JOIN org_department d ON d.id=dv.department_id
        LEFT JOIN org_staff s ON s.id=dv.custodian_id
        WHERE dv.is_deleted=0 AND dv.need_calibration=1 AND dv.status<>'scrapped'
        ORDER BY date(dv.next_calibration_date)
    """)
    data = []
    for r in rows:
        v = list(r.values())
        left = days_between(r["next_calibration_date"])
        if left is None:
            v.append("无需校准")
        elif left < 0:
            v.append(f"已过期 {-left} 天")
        elif left <= CALIBRATION_WARN_DAYS:
            v.append(f"临期（剩 {left} 天）")
        else:
            v.append(f"正常（剩 {left} 天）")
        data.append(v)
    wb = Workbook()
    ws = _sheet(wb, "校准提醒台账",
                ["资产编号", "设备名称", "型号", "使用部门", "存放位置", "责任人",
                 "校准周期(天)", "上次校准", "下次校准", "提醒状态"],
                data, int_cols=(7,), widths=[14, 26, 20, 14, 16, 12, 13, 13, 13, 18], first=True)
    for ri, r in enumerate(data, start=2):
        st = r[9]
        f = EXPIRED_FILL if st.startswith("已过期") else (EXPIRING_FILL if st.startswith("临期") else None)
        if f:
            for ci in range(1, 11):
                ws.cell(row=ri, column=ci).fill = f
    return _stream(wb, _file_name("设备校准提醒"))


@router.get("/borrows.xlsx")
def export_borrows(conn=Depends(get_conn), user: dict = Depends(current_user)):
    rows = query_all(conn, """
        SELECT b.borrow_date, dv.asset_no, dv.name, b.borrower_name,
               COALESCE(d.name,'-') AS dept, b.due_date, b.return_date, b.status, b.purpose
        FROM device_borrow b JOIN device dv ON dv.id=b.device_id
        LEFT JOIN org_department d ON d.id=b.borrower_dept_id
        ORDER BY b.id DESC
    """)
    label = {"borrowed": "借用中", "returned": "已归还", "overdue": "已逾期"}
    data = []
    for r in rows:
        v = list(r.values())
        v[7] = label.get(v[7], v[7])
        data.append(v)
    wb = Workbook()
    _sheet(wb, "设备借用台账",
           ["借出日期", "资产编号", "设备名称", "借用人", "部门", "应还日期", "归还日期", "状态", "用途"],
           data, widths=[13, 14, 26, 12, 14, 13, 13, 10, 26], first=True)
    return _stream(wb, _file_name("设备借用台账"))
